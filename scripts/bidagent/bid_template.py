"""Word 导出模板：把两册标书的排版参数从代码里抽出来，做成可配置、可复用、可校验的模板。

设计约束（很重要）：**默认预设 `standard` 必须与历史输出一致**。
`bidding_documents.build_bid_documents` 原本把 A4 / 2.2-2.5cm 页边距 / 宋体 10.5 /
宋体黑色标题 / PAGE 域页码硬编码在函数里；这些值被原样搬进 `standard`。

关键点：历史代码**故意没有**设置某些属性（标题字号、加粗、段间距、页脚对齐、图片对齐），
它们沿用 Word 文档默认值。所以在模板里用 **`null` 表示「不修改」**，
`standard` 把这些字段全设为 null —— 这样不传模板时的产出与改动前一致（有回归测试盯着）。
其余字段（字体 / 颜色 / 纸张 / 页边距 / 页码格式等）历史代码确实显式设置过，就照抄其值。

模板结构（全部可选，缺省从 `standard` 继承；null = 不修改）：

    page    {paper, orientation, margins_cm{top,bottom,left,right}, first_page_different}
    body    {font, size_pt, line_spacing|null, first_line_indent_chars|null, align|null}
    heading {font, color, size_pts[1..4]|null, bold|null, align|null,
             space_before_pt|null, space_after_pt|null}
    table   {style, header_shading, font_size_pt|null, caption_prefix}
    image   {max_width_cm, max_height_cm, align|null, caption_prefix}
    header  {enabled, text, align|null, size_pt, style(plain|rules|top-bar)}
    footer  {enabled, text, align|null, size_pt, style, page_number_format}

页码用**可刷新的 PAGE 域**（不是写死的数字）；`{page}` 是占位符，例如「第{page}页」。
"""
from __future__ import annotations

import copy
import json
import os
import re
import tempfile
from pathlib import Path
from typing import Any

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .settings import settings


class TemplateError(ValueError):
    """模板模块的可预期错误（对应 HTTP 422）。"""


# 纸张尺寸（厘米，纵向）
PAPER_SIZES: dict[str, tuple[float, float]] = {
    "a4": (21.0, 29.7),
    "a3": (29.7, 42.0),
    "a5": (14.8, 21.0),
    "b5": (17.6, 25.0),
    "letter": (21.59, 27.94),
    "16k": (18.4, 26.0),
}
ORIENTATIONS = ("portrait", "landscape")
ALIGNMENTS = ("left", "center", "right", "justify")
HEADER_FOOTER_STYLES = ("plain", "rules", "top-bar")
TABLE_STYLES = ("Table Grid", "Light Grid Accent 1", "Light List Accent 1", "Medium Shading 1 Accent 1")

_ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
_HEX_COLOR = re.compile(r"^[0-9A-Fa-f]{6}$")
_TEMPLATE_NAME = re.compile(r"^[A-Za-z0-9_-]{1,40}$")

# 与历史硬编码输出逐项对齐；null = 不修改（沿用 Word 文档默认）
DEFAULT_TEMPLATE: dict = {
    "name": "standard",
    "label": "标准投标（默认）",
    "description": "与历史输出一致：A4 纵向、宋体 10.5、宋体黑色标题；标题字号/间距沿用 Word 默认。",
    "page": {"paper": "a4", "orientation": "portrait",
             "margins_cm": {"top": 2.2, "bottom": 2.2, "left": 2.5, "right": 2.5},
             "first_page_different": False},
    "body": {"font": "宋体", "size_pt": 10.5, "line_spacing": None,
             "first_line_indent_chars": None, "align": None},
    "heading": {"font": "宋体", "color": "000000", "size_pts": None, "bold": None,
                "align": None, "space_before_pt": None, "space_after_pt": None},
    "table": {"style": "Table Grid", "header_shading": "", "font_size_pt": None, "caption_prefix": "表"},
    "image": {"max_width_cm": 16.0, "max_height_cm": 22.0, "align": None, "caption_prefix": "图"},
    "header": {"enabled": False, "text": "", "align": None, "size_pt": 9.0, "style": "plain"},
    "footer": {"enabled": True, "text": "", "align": None, "size_pt": 9.0, "style": "plain",
               "page_number_format": "{page}"},
}

_PRESETS: dict[str, dict] = {
    "standard": {},
    "a4-visual": {
        "label": "A4 图文版",
        "description": "A4 纵向、更大字号与行距、蓝色标题、页眉色条、页脚「第N页」——适合强调可读性的技术标。",
        "page": {"margins_cm": {"top": 2.0, "bottom": 2.0, "left": 2.4, "right": 2.4}},
        "body": {"size_pt": 11.0, "line_spacing": 1.5, "first_line_indent_chars": 2, "align": "justify"},
        "heading": {"color": "1A4D8F", "size_pts": [18.0, 15.0, 13.0, 11.5], "bold": True, "align": "left",
                    "space_before_pt": 14.0, "space_after_pt": 8.0},
        "table": {"header_shading": "D9E1F2"},
        "image": {"align": "center"},
        "header": {"enabled": True, "align": "right", "style": "top-bar"},
        "footer": {"align": "center", "style": "rules", "page_number_format": "第{page}页"},
    },
    "a3-landscape": {
        "label": "A3 横版（大图表）",
        "description": "A3 横向、表头浅底、图片可用宽度大——适合含大表格 / 大架构图的方案。",
        "page": {"paper": "a3", "orientation": "landscape",
                 "margins_cm": {"top": 2.0, "bottom": 2.0, "left": 2.0, "right": 2.0}},
        "body": {"size_pt": 10.5, "line_spacing": 1.3, "align": "justify"},
        "heading": {"size_pts": [18.0, 15.0, 13.0, 11.5], "bold": True},
        "table": {"header_shading": "EDEDED", "font_size_pt": 10.0},
        "image": {"max_width_cm": 37.0, "max_height_cm": 26.0, "align": "center"},
        "footer": {"align": "center", "page_number_format": "第{page}页"},
    },
}


# ---------- 校验工具 ----------


def _number(value: Any, low: float, high: float, field: str) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise TemplateError(f"{field} 必须是数字：{value!r}") from exc
    if not low <= number <= high:
        raise TemplateError(f"{field} 超出范围（{low}~{high}）：{number}")
    return number


def _optional_number(value: Any, low: float, high: float, field: str):
    if value is None or value == "":
        return None
    return _number(value, low, high, field)


def _choice(value: Any, allowed, field: str, default: str) -> str:
    text = str(value if value not in (None, "") else default)
    if text not in allowed:
        raise TemplateError(f"{field} 取值非法：{text!r}（可选：{'、'.join(allowed)}）")
    return text


def _optional_choice(value: Any, allowed, field: str):
    if value is None or value == "":
        return None
    return _choice(value, allowed, field, allowed[0])


def _color(value: Any, field: str, default: str = "") -> str:
    text = str(value if value not in (None, "") else default).lstrip("#")
    if not text:
        return ""
    if not _HEX_COLOR.fullmatch(text):
        raise TemplateError(f"{field} 必须是 6 位十六进制颜色：{value!r}")
    return text.upper()


def _merge(base: dict, overlay: dict) -> dict:
    merged = copy.deepcopy(base)
    for key, value in (overlay or {}).items():
        if isinstance(value, dict) and isinstance(merged.get(key), dict):
            merged[key].update(value)
        else:
            merged[key] = value
    return merged


def _preset_template(base: str) -> dict:
    preset = _PRESETS.get(base)
    if preset is None:
        raise TemplateError(f"未知预设：{base}")
    merged = _merge(DEFAULT_TEMPLATE, preset)
    # 预设自身不带 name（只有 label/description），这里补上，否则解析结果会一直显示成 standard。
    merged["name"] = base
    return merged


# ---------- 归一化 ----------


def normalize_template(raw: Any, *, base: str = "standard") -> dict:
    """把（可能不完整的）模板合并到基础模板上并校验，返回完整模板。"""
    if raw in (None, ""):
        raw = {}
    if isinstance(raw, str):
        raise TemplateError("模板应为对象；按名称引用请使用 resolve_template")
    if not isinstance(raw, dict):
        raise TemplateError("模板必须是对象")
    unknown = sorted(set(raw) - set(DEFAULT_TEMPLATE))
    if unknown:
        raise TemplateError("模板含不支持的字段：" + "、".join(unknown)
                            + f"（可用：{'、'.join(sorted(DEFAULT_TEMPLATE))}）")

    merged = _merge(_preset_template(base), raw)

    page = merged["page"]
    page["paper"] = _choice(page.get("paper"), tuple(PAPER_SIZES), "page.paper", "a4")
    page["orientation"] = _choice(page.get("orientation"), ORIENTATIONS, "page.orientation", "portrait")
    margins = page.setdefault("margins_cm", {})
    for side in ("top", "bottom", "left", "right"):
        margins[side] = _number(margins.get(side, 2.2), 0.5, 8.0, f"page.margins_cm.{side}")
    page["first_page_different"] = bool(page.get("first_page_different"))

    body = merged["body"]
    body["font"] = str(body.get("font") or "宋体")[:40]
    body["size_pt"] = _number(body.get("size_pt", 10.5), 6.0, 30.0, "body.size_pt")
    body["line_spacing"] = _optional_number(body.get("line_spacing"), 0.8, 3.0, "body.line_spacing")
    indent = _optional_number(body.get("first_line_indent_chars"), 0, 6, "body.first_line_indent_chars")
    body["first_line_indent_chars"] = None if indent is None else int(indent)
    body["align"] = _optional_choice(body.get("align"), ALIGNMENTS, "body.align")

    heading = merged["heading"]
    heading["font"] = str(heading.get("font") or "宋体")[:40]
    heading["color"] = _color(heading.get("color"), "heading.color", "000000")
    sizes = heading.get("size_pts")
    if sizes is None:
        heading["size_pts"] = None
    else:
        if not isinstance(sizes, (list, tuple)) or len(sizes) != 4:
            raise TemplateError("heading.size_pts 必须是 4 个数字（一~四级标题字号），或 null")
        heading["size_pts"] = [_number(item, 6.0, 30.0, "heading.size_pts") for item in sizes]
    heading["bold"] = None if heading.get("bold") is None else bool(heading.get("bold"))
    heading["align"] = _optional_choice(heading.get("align"), ALIGNMENTS, "heading.align")
    heading["space_before_pt"] = _optional_number(heading.get("space_before_pt"), 0.0, 60.0,
                                                  "heading.space_before_pt")
    heading["space_after_pt"] = _optional_number(heading.get("space_after_pt"), 0.0, 60.0,
                                                 "heading.space_after_pt")

    table = merged["table"]
    table["style"] = _choice(table.get("style"), TABLE_STYLES, "table.style", "Table Grid")
    table["header_shading"] = _color(table.get("header_shading"), "table.header_shading")
    table["font_size_pt"] = _optional_number(table.get("font_size_pt"), 6.0, 24.0, "table.font_size_pt")
    table["caption_prefix"] = str(table.get("caption_prefix") or "表")[:8]

    image = merged["image"]
    image["max_width_cm"] = _number(image.get("max_width_cm", 16.0), 2.0, 60.0, "image.max_width_cm")
    image["max_height_cm"] = _number(image.get("max_height_cm", 22.0), 2.0, 60.0, "image.max_height_cm")
    image["align"] = _optional_choice(image.get("align"), ALIGNMENTS, "image.align")
    image["caption_prefix"] = str(image.get("caption_prefix") or "图")[:8]

    for key in ("header", "footer"):
        block = merged[key]
        block["enabled"] = bool(block.get("enabled"))
        block["text"] = str(block.get("text") or "")[:120]
        block["align"] = _optional_choice(block.get("align"), ALIGNMENTS, f"{key}.align")
        block["size_pt"] = _number(block.get("size_pt", 9.0), 5.0, 20.0, f"{key}.size_pt")
        block["style"] = _choice(block.get("style"), HEADER_FOOTER_STYLES, f"{key}.style", "plain")
    merged["footer"]["page_number_format"] = str(merged["footer"].get("page_number_format") or "{page}")[:40]

    merged["name"] = str(merged.get("name") or base)[:40]
    merged["label"] = str(merged.get("label") or merged["name"])[:60]
    merged["description"] = str(merged.get("description") or "")[:200]
    return merged


def list_presets() -> dict:
    """系统预设（返回已解析后的关键排版参数，便于前端/模型直接看）。"""
    items = []
    for name in _PRESETS:
        resolved = normalize_template({}, base=name)
        items.append({"name": name, "label": resolved["label"], "description": resolved["description"],
                      "paper": resolved["page"]["paper"], "orientation": resolved["page"]["orientation"],
                      "body_font": resolved["body"]["font"], "body_size_pt": resolved["body"]["size_pt"]})
    return {"presets": items, "fields": sorted(DEFAULT_TEMPLATE)}


# ---------- 用户模板存取 ----------


def template_dir() -> Path:
    return Path(settings.bid_state_dir) / "templates"


def _templates_path() -> Path:
    return template_dir() / "templates.json"


def _read_store() -> dict:
    path = _templates_path()
    if not path.is_file():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        raise TemplateError(f"模板库无法解析：{type(exc).__name__}") from exc
    return data if isinstance(data, dict) else {}


def _write_store(store: dict) -> None:
    path = _templates_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(store, ensure_ascii=False, indent=2)
    handle_fd, temporary = tempfile.mkstemp(dir=path.parent, suffix=".tmp")
    try:
        with os.fdopen(handle_fd, "w", encoding="utf-8") as handle:
            handle.write(payload)
        os.replace(temporary, path)
    finally:
        Path(temporary).unlink(missing_ok=True)


def list_templates() -> dict:
    stored = []
    for name, template in sorted(_read_store().items()):
        stored.append({"name": name, "label": template.get("label") or name,
                       "paper": (template.get("page") or {}).get("paper"),
                       "orientation": (template.get("page") or {}).get("orientation")})
    return {"templates": stored, "presets": [item["name"] for item in list_presets()["presets"]]}


def save_template(name: str, template: Any) -> dict:
    if not _TEMPLATE_NAME.fullmatch(str(name or "")):
        raise TemplateError(f"模板名非法（只允许字母数字与 _ -，≤40 字符）：{name!r}")
    if name in _PRESETS:
        raise TemplateError(f"「{name}」是系统预设名，请换一个名字（预设不可覆盖）")
    if not isinstance(template, dict):
        raise TemplateError("模板必须是对象（可含 base 指定继承的预设）")
    base = str(template.get("base") or "standard")
    if base not in _PRESETS:
        raise TemplateError(f"base 必须是预设名：{'、'.join(_PRESETS)}")
    resolved = normalize_template({k: v for k, v in template.items() if k != "base"}, base=base)
    resolved["name"] = name
    store = _read_store()
    store[name] = resolved
    _write_store(store)
    return {"ok": True, "name": name, "template": resolved}


def load_template(name: str) -> dict:
    template = _read_store().get(str(name))
    if template is None:
        raise TemplateError(f"未找到模板：{name}（系统预设：{'、'.join(_PRESETS)}）")
    return normalize_template(template, base="standard")


def delete_template(name: str) -> dict:
    store = _read_store()
    if str(name) not in store:
        raise TemplateError(f"未找到模板：{name}")
    store.pop(str(name))
    _write_store(store)
    return {"ok": True, "deleted": name}


def resolve_template(ref: Any) -> dict:
    """把「预设名 / 用户模板名 / 内联对象 / 空」统一解析为完整模板。"""
    if isinstance(ref, dict):
        base = str(ref.get("base") or "standard")
        if base not in _PRESETS:
            raise TemplateError(f"base 必须是预设名：{'、'.join(_PRESETS)}")
        return normalize_template({k: v for k, v in ref.items() if k != "base"}, base=base)
    if ref in (None, ""):
        return normalize_template({}, base="standard")
    name = str(ref)
    if name in _PRESETS:
        # 注意：预设要靠 base 生效，不能把它当 raw 叠上去 —— 那样 name 会停留成 standard
        return normalize_template({}, base=name)
    return load_template(name)


# ---------- 应用到 docx ----------


def _set_style_font(style, font: str, size_pt, color: str = "") -> None:
    style.font.name = font
    if size_pt is not None:
        style.font.size = Pt(size_pt)
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)


def _strip_paragraph_borders(style) -> None:
    for border in list(style.element.xpath("./w:pPr/w:pBdr")):
        border.getparent().remove(border)


def _paragraph_border(paragraph, edge: str, size: int, color: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    element = OxmlElement(f"w:{edge}")
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), str(size))
    element.set(qn("w:space"), "1")
    element.set(qn("w:color"), color or "000000")
    borders.append(element)


def _paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)


def _apply_band_style(paragraph, style: str, color: str) -> None:
    if style == "rules":
        _paragraph_border(paragraph, "top", 6, color)
        _paragraph_border(paragraph, "bottom", 6, color)
    elif style == "top-bar":
        _paragraph_border(paragraph, "bottom", 18, color)
        _paragraph_shading(paragraph, "F2F6FC")


def add_caption(doc, text: str, *, align: str = "center", size_pt: float | None = None,
                bold: bool = True) -> None:
    """插入图表题注（表题在表上方、图题在图下方，由调用方决定插入位置）。"""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.alignment = _ALIGN_MAP[align]
    run = paragraph.add_run(text)
    if size_pt:
        run.font.size = Pt(size_pt)
    run.font.bold = bool(bold)


def shade_cell(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def style_table(table, template: dict, header_row=None) -> None:
    """按模板给表格上表头底色与字号。`header_row` 默认取第一行（bid_html 会显式传它认定的表头行）。"""
    config = template["table"]
    if config["header_shading"]:
        rows = [header_row] if header_row is not None else list(table.rows[:1])
        for row in rows:
            for cell in row.cells:
                shade_cell(cell, config["header_shading"])
    if config["font_size_pt"]:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(config["font_size_pt"])


def add_figure(doc, path: Path, template: dict, *, caption: str = "") -> None:
    """按模板的图片设置插入一张图（等比缩放到画框内），可选图注。"""
    from PIL import Image, ImageOps

    image_cfg = template["image"]
    with Image.open(path) as image:
        image = ImageOps.exif_transpose(image).convert("RGB")
        width, height = image.size
        scale = min(image_cfg["max_width_cm"] / width, image_cfg["max_height_cm"] / height)
        import io as _io
        stream = _io.BytesIO()
        image.save(stream, format="PNG")
        stream.seek(0)
        paragraph = doc.add_paragraph()
        if image_cfg["align"] is not None:
            paragraph.paragraph_format.alignment = _ALIGN_MAP[image_cfg["align"]]
        paragraph.add_run().add_picture(stream, width=Cm(width * scale), height=Cm(height * scale))
    if caption:
        add_caption(doc, caption)


def apply_document_template(doc, template: dict, *, banner: str = "",
                            banner_color: str = "A04040") -> None:
    """把模板应用到文档：页面、正文/标题样式、页眉页脚、页码域。

    `banner` 是「草稿 DRAFT 待核验」这类强制提示（与模板无关，草稿态必须出现）。
    """
    section = doc.sections[0]
    width, height = PAPER_SIZES[template["page"]["paper"]]
    if template["page"]["orientation"] == "landscape":
        width, height = height, width
    section.page_width, section.page_height = Cm(width), Cm(height)
    margins = template["page"]["margins_cm"]
    section.top_margin, section.bottom_margin = Cm(margins["top"]), Cm(margins["bottom"])
    section.left_margin, section.right_margin = Cm(margins["left"]), Cm(margins["right"])
    section.different_first_page_header_footer = bool(template["page"]["first_page_different"])

    body = template["body"]
    normal = doc.styles["Normal"]
    _set_style_font(normal, body["font"], body["size_pt"])
    _strip_paragraph_borders(normal)
    if body["align"] is not None:
        normal.paragraph_format.alignment = _ALIGN_MAP[body["align"]]
    if body["line_spacing"] is not None:
        normal.paragraph_format.line_spacing = body["line_spacing"]
    if body["first_line_indent_chars"] is not None:
        normal.paragraph_format.first_line_indent = Pt(body["size_pt"] * body["first_line_indent_chars"])

    heading = template["heading"]
    # 注意：`standard`（未给 size_pts）只归一化历史用到的 Title/Subtitle/Heading 1-3。
    # 历史代码的样式列表就到 Heading 3；多动 Heading 4 会让默认产出多出一处 diff。
    # 显式给出 size_pts（4 个数字）的模板才同时作用到 Heading 4。
    levels = range(1, 5) if heading["size_pts"] else range(1, 4)
    for level in levels:
        style = doc.styles[f"Heading {level}"]
        size = heading["size_pts"][level - 1] if heading["size_pts"] else None
        _set_style_font(style, heading["font"], size, heading["color"])
        if heading["bold"] is not None:
            style.font.bold = bool(heading["bold"])
        if heading["align"] is not None:
            style.paragraph_format.alignment = _ALIGN_MAP[heading["align"]]
        if heading["space_before_pt"] is not None:
            style.paragraph_format.space_before = Pt(heading["space_before_pt"])
        if heading["space_after_pt"] is not None:
            style.paragraph_format.space_after = Pt(heading["space_after_pt"])
        _strip_paragraph_borders(style)
    for style_name in ("Title", "Subtitle"):
        style = doc.styles[style_name]
        size = heading["size_pts"][0] if heading["size_pts"] else None
        _set_style_font(style, heading["font"], size, heading["color"])
        _strip_paragraph_borders(style)

    header_cfg = template["header"]
    if header_cfg["enabled"] or banner:
        paragraph = section.header.paragraphs[0]
        paragraph.clear()          # 用 clear() 而不是 .text="" —— 后者会留下一个空的 <w:r/>
        if header_cfg["align"] is not None:
            paragraph.paragraph_format.alignment = _ALIGN_MAP[header_cfg["align"]]
        if header_cfg["text"]:
            run = paragraph.add_run(header_cfg["text"])
            run.font.name = heading["font"]
            run.font.size = Pt(header_cfg["size_pt"])
            run.element.rPr.rFonts.set(qn("w:eastAsia"), heading["font"])
        if banner:
            if header_cfg["text"]:
                paragraph.add_run("    ")
            run = paragraph.add_run(banner)
            run.font.color.rgb = RGBColor.from_string(banner_color)
            run.font.size = Pt(14)
        _apply_band_style(paragraph, header_cfg["style"], heading["color"])

    footer_cfg = template["footer"]
    if footer_cfg["enabled"]:
        paragraph = section.footer.paragraphs[0]
        paragraph.clear()          # 同上：避免留下空的 <w:r/>
        if footer_cfg["align"] is not None:
            paragraph.paragraph_format.alignment = _ALIGN_MAP[footer_cfg["align"]]
        if footer_cfg["text"]:
            paragraph.add_run(footer_cfg["text"]).font.size = Pt(footer_cfg["size_pt"])
        before, _, after = footer_cfg["page_number_format"].partition("{page}")
        if before:
            paragraph.add_run(before).font.size = Pt(footer_cfg["size_pt"])
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        field.set(qn("w:dirty"), "true")
        run_element, text_element = OxmlElement("w:r"), OxmlElement("w:t")
        text_element.text = "待更新"
        run_element.append(text_element)
        field.append(run_element)
        paragraph._p.append(field)
        if after:
            paragraph.add_run(after).font.size = Pt(footer_cfg["size_pt"])
        _apply_band_style(paragraph, footer_cfg["style"], heading["color"])


__all__ = ["TemplateError", "PAPER_SIZES", "ORIENTATIONS", "ALIGNMENTS", "HEADER_FOOTER_STYLES",
           "TABLE_STYLES", "DEFAULT_TEMPLATE", "normalize_template", "list_presets",
           "list_templates", "save_template", "load_template", "delete_template",
           "resolve_template", "apply_document_template", "template_dir",
           "add_caption", "add_figure", "shade_cell", "style_table"]
